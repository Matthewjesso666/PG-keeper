"""Document import and text extraction helpers."""

import logging
import shutil
from datetime import datetime
from pathlib import Path
from typing import Optional, Tuple
from uuid import uuid4

from docx import Document as DocxDocument
from pypdf import PdfReader

from wcb_app.models import CaseState, Document, ExtractionResult
from wcb_app.storage import save_case_state
from wcb_app.utils import ensure_dir

logger = logging.getLogger(__name__)

SUPPORTED_TYPES = {".pdf", ".docx", ".txt"}
DOCUMENT_TYPES = [
    "decision",
    "medical",
    "correspondence",
    "employer",
    "wage",
    "rtw",
    "other",
]


def import_document(case_state: CaseState, source_path: Path, doc_type: Optional[str] = None) -> Tuple[Optional[Document], str]:
    """Copy a document into the case and extract text."""
    if not source_path.exists():
        return None, "File not found. Please check the path and try again."

    extension = source_path.suffix.lower()
    if extension not in SUPPORTED_TYPES:
        return None, "Unsupported file type. Please use PDF, DOCX, or TXT."

    doc_type_value = doc_type or "other"
    stored_name = f"{datetime.utcnow().strftime('%Y%m%d_%H%M%S')}_{source_path.name}"
    dest_path = Path(case_state.base_path) / "imports" / stored_name
    ensure_dir(dest_path.parent)
    shutil.copy2(source_path, dest_path)

    document = Document(
        filename=source_path.name,
        doc_type=doc_type_value,
        original_path=str(source_path),
        stored_path=str(dest_path),
    )

    extraction_result = extract_text(document, case_state)
    if extraction_result.success:
        document.extracted_text_path = extraction_result.text_path
        document.status = "text extracted"
    else:
        document.status = "imported"

    case_state.documents.append(document)
    save_case_state(case_state)
    logger.info("Imported document %s", document.filename)

    message = (
        "Imported and extracted text."
        if extraction_result.success
        else f"Imported with warnings: {extraction_result.error_message}"
    )
    return document, message


def extract_text(document: Document, case_state: CaseState) -> ExtractionResult:
    """Extract text for a given document."""
    stored_path = Path(document.stored_path)
    extracted_dir = Path(case_state.base_path) / "notes" / "extracted"
    ensure_dir(extracted_dir)
    text_path = extracted_dir / f"{uuid4().hex}.txt"

    try:
        if stored_path.suffix.lower() == ".pdf":
            text = _read_pdf(stored_path)
        elif stored_path.suffix.lower() == ".docx":
            text = _read_docx(stored_path)
        else:
            text = stored_path.read_text(encoding="utf-8", errors="ignore")
        text_path.write_text(text.strip(), encoding="utf-8")
        return ExtractionResult(doc_id=document.id, success=True, text_path=str(text_path))
    except Exception as exc:  # noqa: BLE001
        logger.exception("Failed to extract text from %s", stored_path)
        return ExtractionResult(
            doc_id=document.id,
            success=False,
            text_path=None,
            error_message=f"Could not extract text: {exc}",
        )


def _read_pdf(path: Path) -> str:
    """Extract text from a PDF using pypdf."""
    reader = PdfReader(str(path))
    texts = []
    for page in reader.pages:
        try:
            texts.append(page.extract_text() or "")
        except Exception as exc:  # noqa: BLE001
            logger.warning("Skipped a PDF page due to error: %s", exc)
            texts.append("[Page could not be read]")
    return "\n".join(texts)


def _read_docx(path: Path) -> str:
    """Extract text from a DOCX file."""
    doc = DocxDocument(path)
    paragraphs = [p.text for p in doc.paragraphs if p.text]
    return "\n".join(paragraphs)
